#!/usr/bin/env python3
"""Fase 4 — Relatórios Agendados de Observabilidade.

Gera relatórios diários ou semanais em Excel (.xlsx) e Markdown (.md)
para três audiências: COI/IOC, SOC e iPaaS.

Uso:
    python scripts/generate_reports.py --period daily  --output-dir reports/
    python scripts/generate_reports.py --period weekly --output-dir reports/

Variáveis de ambiente:
    DATABASE_URL  — PostgreSQL asyncpg ou psycopg2 (obrigatória)
    REPORT_PERIOD — daily | weekly (padrão: daily; sobreposto por --period)
    REPORT_OUTPUT — diretório de saída (padrão: reports/)

Dependências extras:
    pip install -e ".[reports]"
    # openpyxl>=3.1, jinja2>=3.1, psycopg2-binary>=2.9
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import psycopg2
import psycopg2.extras
from jinja2 import Environment, BaseLoader
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%dT%H:%M:%S",
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Conexão
# ---------------------------------------------------------------------------

def _sync_url(raw: str) -> str:
    """Converte URL asyncpg → psycopg2."""
    url = re.sub(r"^postgresql\+asyncpg://", "postgresql://", raw)
    url = re.sub(r"^postgres://", "postgresql://", url)
    return url


def _connect(database_url: str) -> psycopg2.extensions.connection:
    url = _sync_url(database_url)
    conn = psycopg2.connect(url, cursor_factory=psycopg2.extras.RealDictCursor)
    conn.set_session(readonly=True, autocommit=True)
    return conn


# ---------------------------------------------------------------------------
# Janela de tempo
# ---------------------------------------------------------------------------

def _time_window(period: str) -> tuple[datetime, datetime]:
    now = datetime.now(timezone.utc)
    if period == "weekly":
        since = now - timedelta(days=7)
    else:  # daily
        since = now - timedelta(days=1)
    return since, now


# ---------------------------------------------------------------------------
# Queries
# ---------------------------------------------------------------------------

_Q_COI_SUMMARY = """
SELECT
    COUNT(*)                                                          AS total,
    COUNT(*) FILTER (WHERE diagnosis_correct IS NOT NULL)            AS verified,
    COUNT(*) FILTER (WHERE diagnosis_correct = true)                 AS correct,
    ROUND(AVG(latency_ms))                                           AS avg_latency_ms,
    ROUND(PERCENTILE_CONT(0.5) WITHIN GROUP (ORDER BY latency_ms))  AS p50_latency_ms,
    ROUND(PERCENTILE_CONT(0.95) WITHIN GROUP (ORDER BY latency_ms)) AS p95_latency_ms,
    COUNT(DISTINCT agent_domain)                                      AS domains
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s;
"""

_Q_COI_BY_DOMAIN = """
SELECT
    COALESCE(agent_domain, 'n/a')                                    AS domain,
    COUNT(*)                                                          AS total,
    COUNT(*) FILTER (WHERE diagnosis_correct = true)                  AS correct,
    ROUND(AVG(model_confidence)::numeric, 3)                          AS avg_confidence,
    ROUND(AVG(latency_ms))                                            AS avg_latency_ms
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
GROUP BY 1
ORDER BY 2 DESC;
"""

_Q_COI_TOP_CAUSES = """
SELECT
    LEFT(probable_root_cause, 100) AS causa,
    COUNT(*)                        AS ocorrencias
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
  AND probable_root_cause IS NOT NULL
GROUP BY 1
ORDER BY 2 DESC
LIMIT 10;
"""

_Q_COI_RECENT = """
SELECT
    id::text,
    created_at,
    interface_type,
    connector_source_system,
    LEFT(probable_root_cause, 120) AS causa,
    model_confidence,
    latency_ms,
    CASE WHEN diagnosis_correct IS NULL THEN 'Pendente'
         WHEN diagnosis_correct THEN 'Correto'
         ELSE 'Incorreto' END       AS verificacao
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
ORDER BY created_at DESC
LIMIT 50;
"""

_Q_SOC_SUMMARY = """
SELECT
    COUNT(*)                                                            AS total,
    COUNT(*) FILTER (WHERE pii_detected = true)                        AS pii_total,
    COUNT(*) FILTER (WHERE pii_detected = true AND redaction_applied)  AS pii_redigido,
    COUNT(*) FILTER (WHERE pii_detected = true AND NOT redaction_applied) AS pii_nao_redigido,
    COUNT(*) FILTER (WHERE sensitivity_level = 'critical')             AS critical_total,
    COUNT(*) FILTER (WHERE sensitivity_level = 'high')                 AS high_total,
    COUNT(DISTINCT agent_domain) FILTER (WHERE pii_detected = true)    AS dominios_pii
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s;
"""

_Q_SOC_BY_SENSITIVITY = """
SELECT
    COALESCE(sensitivity_level, 'não classificado') AS nivel,
    COUNT(*)                                          AS total,
    COUNT(*) FILTER (WHERE pii_detected = true)       AS com_pii,
    COUNT(*) FILTER (WHERE redaction_applied = true)  AS redigidos
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
GROUP BY 1
ORDER BY 2 DESC;
"""

_Q_SOC_PII_INCIDENTS = """
SELECT
    LEFT(id::text, 8)              AS id,
    created_at,
    interface_type,
    agent_domain,
    sensitivity_level,
    pii_detected,
    redaction_applied,
    CASE WHEN diagnosis_correct IS NULL THEN 'Pendente'
         WHEN diagnosis_correct THEN 'Correto'
         ELSE 'Incorreto' END       AS verificacao
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
  AND pii_detected = true
ORDER BY created_at DESC
LIMIT 50;
"""

_Q_IPAAS_SUMMARY = """
SELECT
    COUNT(*)                                                          AS total,
    COUNT(DISTINCT interface_type)                                    AS interfaces_unicas,
    COUNT(DISTINCT llm_provider_used)                                 AS providers_usados,
    ROUND(AVG(model_confidence)::numeric, 3)                          AS avg_confidence,
    ROUND(PERCENTILE_CONT(0.95) WITHIN GROUP (ORDER BY latency_ms))  AS p95_latency_ms,
    COUNT(*) FILTER (WHERE llm_provider_used LIKE '%%fallback%%')    AS fallbacks_llm,
    COUNT(*) FILTER (WHERE evidence_strength IN ('high','critical'))  AS high_critical_evidence
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s;
"""

_Q_IPAAS_BY_INTERFACE = """
SELECT
    COALESCE(interface_type, 'n/a')                                   AS interface,
    COUNT(*)                                                           AS total,
    ROUND(PERCENTILE_CONT(0.5) WITHIN GROUP (ORDER BY latency_ms))   AS p50_ms,
    ROUND(PERCENTILE_CONT(0.95) WITHIN GROUP (ORDER BY latency_ms))  AS p95_ms,
    ROUND(AVG(model_confidence)::numeric, 3)                           AS avg_confidence
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
GROUP BY 1
ORDER BY 2 DESC
LIMIT 15;
"""

_Q_IPAAS_BY_PROVIDER = """
SELECT
    COALESCE(llm_provider_used, 'n/a')                               AS provider,
    COUNT(*)                                                          AS total,
    ROUND(AVG(latency_ms))                                           AS avg_latency_ms,
    ROUND(AVG(model_confidence)::numeric, 3)                          AS avg_confidence
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
GROUP BY 1
ORDER BY 2 DESC;
"""

_Q_IPAAS_RECENT = """
SELECT
    LEFT(id::text, 8)              AS id,
    created_at,
    interface_type,
    connector_source_system,
    llm_provider_used,
    latency_ms,
    model_confidence,
    evidence_strength
FROM incidents
WHERE created_at BETWEEN %(since)s AND %(until)s
ORDER BY created_at DESC
LIMIT 50;
"""


def _fetch(conn: psycopg2.extensions.connection, sql: str, params: dict) -> list[dict]:
    with conn.cursor() as cur:
        cur.execute(sql, params)
        return [dict(row) for row in cur.fetchall()]


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT  = Font(bold=True, size=13, color="1F4E79")
_ALT_FILL    = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")


def _write_sheet_table(
    ws,
    title: str,
    rows: list[dict],
    start_row: int = 1,
) -> int:
    """Escreve título + tabela formatada; retorna a próxima linha livre."""
    if not rows:
        ws.cell(start_row, 1, f"{title} — sem dados no período").font = Font(italic=True)
        return start_row + 2

    ws.cell(start_row, 1, title).font = _TITLE_FONT
    header_row = start_row + 1
    headers = list(rows[0].keys())

    for col, h in enumerate(headers, 1):
        cell = ws.cell(header_row, col, str(h).replace("_", " ").title())
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(rows, header_row + 1):
        fill = _ALT_FILL if r_idx % 2 == 0 else None
        for col, key in enumerate(headers, 1):
            val = row[key]
            if isinstance(val, datetime):
                val = val.strftime("%Y-%m-%d %H:%M")
            cell = ws.cell(r_idx, col, val)
            if fill:
                cell.fill = fill

    # Auto-largura
    for col in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(r, col).value or ""))
            for r in range(header_row, header_row + len(rows) + 1)
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 4, 50)

    return header_row + len(rows) + 2


def _kv_sheet(ws, title: str, data: dict) -> None:
    ws.cell(1, 1, title).font = _TITLE_FONT
    for r, (k, v) in enumerate(data.items(), 3):
        label = str(k).replace("_", " ").title()
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, v if not isinstance(v, datetime) else v.strftime("%Y-%m-%d %H:%M"))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25


# ---------------------------------------------------------------------------
# Markdown template
# ---------------------------------------------------------------------------

_MD_TEMPLATE = """\
# {{ title }}

**Período:** {{ since }} → {{ until }}  
**Gerado em:** {{ generated_at }}

---

## Resumo Executivo

| Métrica | Valor |
|---|---|
{% for k, v in summary.items() -%}
| {{ k | replace('_', ' ') | title }} | {{ v }} |
{% endfor %}

{% for section in sections %}
## {{ section.title }}

{% if section.rows %}
| {{ section.rows[0].keys() | join(' | ') }} |
| {{ ['---'] * (section.rows[0].keys() | list | length) | join(' | ') }} |
{% for row in section.rows -%}
| {{ row.values() | join(' | ') }} |
{% endfor %}
{% else %}
_Sem dados no período._
{% endif %}

{% endfor %}
---
*IIC — Integration Incident Copilot | Relatório automático Fase 4*
"""


def _render_markdown(
    title: str,
    since: datetime,
    until: datetime,
    summary: dict,
    sections: list[dict],
) -> str:
    env = Environment(loader=BaseLoader(), autoescape=False)
    tpl = env.from_string(_MD_TEMPLATE)
    return tpl.render(
        title=title,
        since=since.strftime("%Y-%m-%d %H:%M UTC"),
        until=until.strftime("%Y-%m-%d %H:%M UTC"),
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        summary=summary,
        sections=sections,
    )


# ---------------------------------------------------------------------------
# Geradores por audiência
# ---------------------------------------------------------------------------

def _report_coi(conn, params: dict, since: datetime, until: datetime, wb: Workbook, out_dir: Path, label: str) -> None:
    summary_rows = _fetch(conn, _Q_COI_SUMMARY, params)
    by_domain    = _fetch(conn, _Q_COI_BY_DOMAIN, params)
    top_causes   = _fetch(conn, _Q_COI_TOP_CAUSES, params)
    recent       = _fetch(conn, _Q_COI_RECENT, params)

    summary = summary_rows[0] if summary_rows else {}
    # Taxa de verificação / acuracidade
    total    = summary.get("total") or 0
    verified = summary.get("verified") or 0
    correct  = summary.get("correct") or 0
    summary["taxa_verificacao_pct"] = f"{100 * verified / total:.1f}%" if total else "n/a"
    summary["acuracidade_pct"]      = f"{100 * correct / verified:.1f}%" if verified else "n/a"

    # Excel
    ws = wb.create_sheet("COI-IOC Resumo")
    _kv_sheet(ws, "COI/IOC — Resumo Executivo", summary)

    ws2 = wb.create_sheet("COI-IOC Por Domínio")
    _write_sheet_table(ws2, "Incidentes por Domínio de Agente", by_domain)

    ws3 = wb.create_sheet("COI-IOC Top Causas")
    _write_sheet_table(ws3, "Top 10 Causas Prováveis", top_causes)

    ws4 = wb.create_sheet("COI-IOC Recentes")
    _write_sheet_table(ws4, "Últimos 50 Incidentes", recent)

    # Markdown
    md = _render_markdown(
        title=f"IIC — COI/IOC: Relatório {label.title()}",
        since=since, until=until,
        summary=summary,
        sections=[
            {"title": "Por Domínio de Agente", "rows": by_domain},
            {"title": "Top 10 Causas Prováveis", "rows": top_causes},
        ],
    )
    (out_dir / f"coi_ioc_{label}.md").write_text(md, encoding="utf-8")
    log.info("COI/IOC markdown → %s", out_dir / f"coi_ioc_{label}.md")


def _report_soc(conn, params: dict, since: datetime, until: datetime, wb: Workbook, out_dir: Path, label: str) -> None:
    summary_rows    = _fetch(conn, _Q_SOC_SUMMARY, params)
    by_sensitivity  = _fetch(conn, _Q_SOC_BY_SENSITIVITY, params)
    pii_incidents   = _fetch(conn, _Q_SOC_PII_INCIDENTS, params)

    summary = summary_rows[0] if summary_rows else {}
    pii_total   = summary.get("pii_total") or 0
    pii_redigido = summary.get("pii_redigido") or 0
    summary["taxa_redacao_pct"] = f"{100 * pii_redigido / pii_total:.1f}%" if pii_total else "n/a"

    ws  = wb.create_sheet("SOC Resumo")
    _kv_sheet(ws, "SOC — Resumo Executivo", summary)

    ws2 = wb.create_sheet("SOC Por Sensibilidade")
    _write_sheet_table(ws2, "Distribuição por Nível de Sensibilidade", by_sensitivity)

    ws3 = wb.create_sheet("SOC Incidentes PII")
    _write_sheet_table(ws3, "Incidentes com PII Detectado (últimos 50)", pii_incidents)

    md = _render_markdown(
        title=f"IIC — SOC: Relatório {label.title()}",
        since=since, until=until,
        summary=summary,
        sections=[
            {"title": "Por Nível de Sensibilidade", "rows": by_sensitivity},
            {"title": "Incidentes com PII (últimos 50)", "rows": pii_incidents},
        ],
    )
    (out_dir / f"soc_{label}.md").write_text(md, encoding="utf-8")
    log.info("SOC markdown → %s", out_dir / f"soc_{label}.md")


def _report_ipaas(conn, params: dict, since: datetime, until: datetime, wb: Workbook, out_dir: Path, label: str) -> None:
    summary_rows  = _fetch(conn, _Q_IPAAS_SUMMARY, params)
    by_interface  = _fetch(conn, _Q_IPAAS_BY_INTERFACE, params)
    by_provider   = _fetch(conn, _Q_IPAAS_BY_PROVIDER, params)
    recent        = _fetch(conn, _Q_IPAAS_RECENT, params)

    summary = summary_rows[0] if summary_rows else {}

    ws  = wb.create_sheet("iPaaS Resumo")
    _kv_sheet(ws, "iPaaS — Resumo Executivo", summary)

    ws2 = wb.create_sheet("iPaaS Por Interface")
    _write_sheet_table(ws2, "Latência e Volume por Tipo de Interface", by_interface)

    ws3 = wb.create_sheet("iPaaS Por LLM Provider")
    _write_sheet_table(ws3, "Diagnósticos por LLM Provider", by_provider)

    ws4 = wb.create_sheet("iPaaS Recentes")
    _write_sheet_table(ws4, "Últimos 50 Incidentes — Visão iPaaS", recent)

    md = _render_markdown(
        title=f"IIC — iPaaS: Relatório {label.title()}",
        since=since, until=until,
        summary=summary,
        sections=[
            {"title": "Por Tipo de Interface", "rows": by_interface},
            {"title": "Por LLM Provider", "rows": by_provider},
        ],
    )
    (out_dir / f"ipaas_{label}.md").write_text(md, encoding="utf-8")
    log.info("iPaaS markdown → %s", out_dir / f"ipaas_{label}.md")


# ---------------------------------------------------------------------------
# Entry-point
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="IIC — gerador de relatórios Fase 4")
    parser.add_argument(
        "--period", choices=["daily", "weekly"],
        default=os.getenv("REPORT_PERIOD", "daily"),
        help="Janela de tempo do relatório (padrão: daily)",
    )
    parser.add_argument(
        "--output-dir",
        default=os.getenv("REPORT_OUTPUT", "reports"),
        help="Diretório de saída (padrão: reports/)",
    )
    parser.add_argument(
        "--audience", choices=["all", "coi", "soc", "ipaas"],
        default="all",
        help="Audiência do relatório (padrão: all)",
    )
    args = parser.parse_args()

    database_url = os.getenv("DATABASE_URL", "")
    if not database_url:
        log.error("DATABASE_URL não definida — abortando.")
        sys.exit(1)

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    since, until = _time_window(args.period)
    label = args.period  # "daily" | "weekly"
    ts    = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
    params: dict[str, Any] = {"since": since, "until": until}

    log.info("Período: %s → %s (%s)", since.isoformat(), until.isoformat(), label)

    conn = _connect(database_url)
    wb   = Workbook()
    wb.remove(wb.active)  # remove aba default em branco

    try:
        audiences = ["coi", "soc", "ipaas"] if args.audience == "all" else [args.audience]

        if "coi" in audiences:
            log.info("Gerando relatório COI/IOC …")
            _report_coi(conn, params, since, until, wb, out_dir, label)

        if "soc" in audiences:
            log.info("Gerando relatório SOC …")
            _report_soc(conn, params, since, until, wb, out_dir, label)

        if "ipaas" in audiences:
            log.info("Gerando relatório iPaaS …")
            _report_ipaas(conn, params, since, until, wb, out_dir, label)

        xlsx_path = out_dir / f"iic_report_{label}_{ts}.xlsx"
        wb.save(xlsx_path)
        log.info("Excel → %s", xlsx_path)

    finally:
        conn.close()

    log.info("Relatórios concluídos em %s/", out_dir)


if __name__ == "__main__":
    main()
